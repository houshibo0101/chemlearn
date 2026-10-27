"""
题目路由：管理员管理题库 + 学生刷题
"""
from datetime import datetime, timezone
from fastapi import APIRouter, Depends, HTTPException, Query, UploadFile, File
from sqlalchemy import select, func, or_
from sqlalchemy.ext.asyncio import AsyncSession
from app.core.database import get_db
from app.core.deps import get_current_user, get_current_admin, get_current_teacher_or_admin
from app.models import User, Question, WrongQuestion, PracticeRecord
from app.schemas import (
    QuestionCreate, QuestionUpdate, QuestionOut,
    QuestionPracticeOut, AnswerSubmit, AnswerResult,
)
from app.services.cache_service import cache

router = APIRouter(prefix="/questions", tags=["题目"])


# ========== 管理员：题库管理 ==========

@router.post("", response_model=QuestionOut, summary="【管理员】上传题目")
async def create_question(
    data: QuestionCreate,
    admin: User = Depends(get_current_teacher_or_admin),
    db: AsyncSession = Depends(get_db),
):
    # 自动生成唯一编号 Q0001
    max_code = (await db.execute(select(Question.question_code).where(Question.question_code.isnot(None)).order_by(Question.id.desc()).limit(1))).scalar()
    next_num = 1
    if max_code and max_code.startswith("Q"):
        try:
            next_num = int(max_code[1:]) + 1
        except ValueError:
            pass
    question_code = f"Q{next_num:04d}"

    # options存为JSON字符串
    options_str = data.options
    if isinstance(options_str, (list, dict)):
        import json as _json
        options_str = _json.dumps(options_str, ensure_ascii=False)

    q = Question(
        question_code=question_code,
        title=data.title,
        content=data.content,
        question_type=data.question_type,
        difficulty=data.difficulty,
        grade=data.grade,
        knowledge_point=data.knowledge_point,
        options=options_str,
        answer=data.answer,
        analysis=data.analysis,
        creator_id=admin.id,
    )
    db.add(q)
    await db.commit()
    await db.refresh(q)
    # 清除缓存
    cache.delete("stats:knowledge_points")
    cache.delete("stats:question_filters")
    cache.delete_pattern("q:list:*")
    return QuestionOut.model_validate(q)


@router.get("/admin", summary="【管理员】获取题目列表（含答案）")
async def list_questions_admin(
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
    keyword: str | None = None,
    grade: str | None = None,
    question_type: str | None = None,
    difficulty: str | None = None,
    admin: User = Depends(get_current_teacher_or_admin),
    db: AsyncSession = Depends(get_db),
):
    query = select(Question)
    count_query = select(func.count(Question.id))
    if keyword:
        # 支持id搜索（纯数字时同时匹配id）
        filters = [
            Question.title.contains(keyword),
            Question.knowledge_point.contains(keyword),
            Question.question_code.contains(keyword),
            Question.content.contains(keyword),
        ]
        # 如果关键词是纯数字，也搜索id
        if keyword.isdigit():
            filters.append(Question.id == int(keyword))
        filter_ = or_(*filters)
        query = query.where(filter_)
        count_query = count_query.where(filter_)
    if grade:
        query = query.where(Question.grade == grade)
        count_query = count_query.where(Question.grade == grade)
    if question_type:
        query = query.where(Question.question_type == question_type)
        count_query = count_query.where(Question.question_type == question_type)
    if difficulty:
        query = query.where(Question.difficulty == difficulty)
        count_query = count_query.where(Question.difficulty == difficulty)

    total = (await db.execute(count_query)).scalar()
    query = query.order_by(Question.id.desc()).offset((page - 1) * page_size).limit(page_size)
    items = (await db.execute(query)).scalars().all()
    return {"total": total, "page": page, "page_size": page_size, "items": [QuestionOut.model_validate(i) for i in items]}


@router.get("/{question_id}", response_model=QuestionOut, summary="【管理员】获取题目详情（含答案）")
async def get_question_admin(
    question_id: int,
    admin: User = Depends(get_current_teacher_or_admin),
    db: AsyncSession = Depends(get_db),
):
    q = (await db.execute(select(Question).where(Question.id == question_id))).scalar_one_or_none()
    if not q:
        raise HTTPException(status_code=404, detail="题目不存在")
    return QuestionOut.model_validate(q)


@router.put("/{question_id}", response_model=QuestionOut, summary="【管理员】更新题目")
async def update_question(
    question_id: int,
    data: QuestionUpdate,
    admin: User = Depends(get_current_teacher_or_admin),
    db: AsyncSession = Depends(get_db),
):
    q = (await db.execute(select(Question).where(Question.id == question_id))).scalar_one_or_none()
    if not q:
        raise HTTPException(status_code=404, detail="题目不存在")
    update_data = data.model_dump(exclude_unset=True)
    # options存为JSON字符串
    if 'options' in update_data and isinstance(update_data['options'], (list, dict)):
        import json as _json
        update_data['options'] = _json.dumps(update_data['options'], ensure_ascii=False)
    for k, v in update_data.items():
        setattr(q, k, v)
    await db.commit()
    await db.refresh(q)
    # 清除缓存
    cache.delete("stats:knowledge_points")
    cache.delete("stats:question_filters")
    cache.delete_pattern("q:list:*")
    cache.delete(f"q:detail:{question_id}")
    return QuestionOut.model_validate(q)


@router.delete("/{question_id}", summary="【管理员】删除题目")
async def delete_question(
    question_id: int,
    admin: User = Depends(get_current_teacher_or_admin),
    db: AsyncSession = Depends(get_db),
):
    q = (await db.execute(select(Question).where(Question.id == question_id))).scalar_one_or_none()
    if not q:
        raise HTTPException(status_code=404, detail="题目不存在")
    await db.delete(q)
    await db.commit()
    # 清除缓存
    cache.delete("stats:knowledge_points")
    cache.delete("stats:question_filters")
    cache.delete_pattern("q:list:*")
    cache.delete(f"q:detail:{question_id}")
    return {"message": "删除成功"}


# ========== 学生：刷题 ==========

@router.get("", summary="获取刷题题目列表（不含答案）")
async def list_questions_practice(
    page: int = Query(1, ge=1),
    page_size: int = Query(10, ge=1, le=100),
    grade: str | None = None,
    question_type: str | None = None,
    difficulty: str | None = None,
    knowledge_point: str | None = None,
    random: bool = Query(False, description="随机排序"),
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    # 随机排序不缓存（每次结果不同）
    if not random:
        cache_key = f"q:list:{page}:{page_size}:{grade}:{question_type}:{difficulty}:{knowledge_point}"
        cached = cache.get_json(cache_key)
        if cached and isinstance(cached, dict) and cached.get("items"):
            # 缓存数据校验：确保第一条有id字段，防止旧格式缓存
            first = cached["items"][0] if cached["items"] else None
            if first and isinstance(first, dict) and "id" in first and "title" in first:
                return cached
            # 格式不对，删除坏缓存
            cache.delete(cache_key)

    query = select(Question)
    count_query = select(func.count(Question.id))
    if grade:
        query = query.where(Question.grade == grade)
        count_query = count_query.where(Question.grade == grade)
    if question_type:
        query = query.where(Question.question_type == question_type)
        count_query = count_query.where(Question.question_type == question_type)
    if difficulty:
        query = query.where(Question.difficulty == difficulty)
        count_query = count_query.where(Question.difficulty == difficulty)
    if knowledge_point:
        query = query.where(Question.knowledge_point.contains(knowledge_point))
        count_query = count_query.where(Question.knowledge_point.contains(knowledge_point))

    total = (await db.execute(count_query)).scalar()
    if random:
        query = query.order_by(func.random())
    else:
        query = query.order_by(Question.id.desc())
    query = query.offset((page - 1) * page_size).limit(page_size)
    items = (await db.execute(query)).scalars().all()
    result = {"total": total, "page": page, "page_size": page_size, "items": [QuestionPracticeOut.model_validate(i) for i in items]}

    # 非随机结果缓存5分钟
    if not random:
        cache.set_json(cache_key, result, expire=300)
    return result


@router.get("/practice/{question_id}", response_model=QuestionPracticeOut, summary="获取单题（不含答案）")
async def get_question_practice(
    question_id: int,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    cache_key = f"q:detail:{question_id}"
    cached = cache.get_json(cache_key)
    if cached:
        return cached

    q = (await db.execute(select(Question).where(Question.id == question_id))).scalar_one_or_none()
    if not q:
        raise HTTPException(status_code=404, detail="题目不存在")
    result = QuestionPracticeOut.model_validate(q)
    cache.set_json(cache_key, result.model_dump(), expire=600)  # 缓存10分钟
    return result


@router.post("/answer", response_model=AnswerResult, summary="提交答案并判分")
async def submit_answer(
    data: AnswerSubmit,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    q = (await db.execute(select(Question).where(Question.id == data.question_id))).scalar_one_or_none()
    if not q:
        raise HTTPException(status_code=404, detail="题目不存在")

    # 判分
    user_ans = data.user_answer.strip()
    correct_ans = q.answer.strip()
    q_type = q.question_type.value if hasattr(q.question_type, 'value') else q.question_type

    is_correct = False
    score = None
    ai_comment = None
    sub_results = None
    answer_total = None  # 该题满分（做题记录用）

    import json as _json

    # 检查是否有分小问
    sub_questions = []
    if q.sub_questions:
        try:
            sub_questions = _json.loads(q.sub_questions)
        except:
            pass

    if sub_questions:
        # 分小问判分
        # 用户答案格式：{"1": "答案1", "2": "答案2"} 或纯文本
        user_sub_answers = {}
        try:
            user_sub_answers = _json.loads(user_ans)
        except:
            # 不是JSON，把整个答案作为第一小问的答案
            user_sub_answers = {str(sub_questions[0].get('id', 1)): user_ans}

        total_score = 0
        got_score = 0
        sub_results = []
        for sq in sub_questions:
            sq_id = str(sq.get('id', ''))
            sq_content = sq.get('content', '')
            sq_answer = sq.get('answer', '')
            sq_score = sq.get('score', 1)
            sq_keywords = sq.get('keywords', [])

            total_score += sq_score
            user_sub_ans = user_sub_answers.get(sq_id, '').strip()

            sub_got = 0
            if sq_keywords:
                for kw in sq_keywords:
                    word = kw.get('word', '')
                    kw_score = kw.get('score', 1)
                    if word and word in user_sub_ans:
                        sub_got += kw_score
            else:
                if user_sub_ans and user_sub_ans.lower() == sq_answer.lower():
                    sub_got = sq_score

            # 小问得分不超过该小问满分
            sub_got = min(sub_got, sq_score)
            got_score += sub_got

            sub_results.append({
                'id': sq_id,
                'content': sq_content,
                'user_answer': user_sub_ans,
                'correct_answer': sq_answer,
                'score': sub_got,
                'full_score': sq_score,
                'is_correct': sub_got >= sq_score * 0.6,
            })

        score = got_score
        is_correct = got_score >= total_score * 0.6
        ai_comment = f"得分：{got_score}/{total_score}"

    elif q_type == 'multiple':
        # 多选题：拆分比较
        user_set = set(user_ans.upper().replace('，', ',').split(','))
        correct_set = set(correct_ans.upper().replace('，', ',').split(','))
        is_correct = user_set == correct_set
    elif q_type == 'subjective' or q_type == 'calculation':
        # 简答题/计算题：关键词匹配
        keywords = []
        if q.keywords:
            try:
                keywords = _json.loads(q.keywords)
            except:
                pass
        if keywords:
            total_score = sum(kw.get('score', 1) for kw in keywords)
            answer_total = total_score
            got_score = 0
            for kw in keywords:
                word = kw.get('word', '')
                s = kw.get('score', 1)
                if word and word in user_ans:
                    got_score += s
            score = got_score
            is_correct = got_score >= total_score * 0.6
        else:
            is_correct = user_ans.lower() == correct_ans.lower()
    else:
        # 单选/填空：简单字符串比较
        is_correct = user_ans.lower() == correct_ans.lower()

    recorded = False
    if not is_correct:
        # 错题去重：同一用户同一题目只保留一条，已存在则次数+1
        existing = (await db.execute(
            select(WrongQuestion).where(
                WrongQuestion.user_id == current_user.id,
                WrongQuestion.question_id == q.id,
            )
        )).scalar_one_or_none()

        if existing:
            existing.wrong_count += 1
            existing.user_answer = data.user_answer
            existing.last_wrong_at = datetime.now(timezone.utc)
            recorded = False
        else:
            wq = WrongQuestion(
                user_id=current_user.id,
                question_id=q.id,
                wrong_count=1,
                user_answer=data.user_answer,
            )
            db.add(wq)
            recorded = True

    # ===== 做题记录：无论对错都留痕（学习报告/做题历史数据源） =====
    # 得分推导：分小问/关键词题用判分结果，客观题对 1 错 0
    if sub_results:
        pr_score = int(score or 0)
        pr_total = sum(int(r.get('full_score', 1) or 1) for r in sub_results)
    elif score is not None and answer_total is not None:
        pr_score, pr_total = int(score), int(answer_total)
    elif score is not None:
        pr_score, pr_total = int(score), max(int(score), 1)
    else:
        pr_score, pr_total = (1, 1) if is_correct else (0, 1)

    db.add(PracticeRecord(
        user_id=current_user.id,
        question_id=q.id,
        source=(data.source or "practice")[:16],
        user_answer=data.user_answer,
        is_correct=bool(is_correct),
        score=pr_score,
        total_score=pr_total,
        time_spent=data.time_spent,
    ))
    await db.commit()

    return AnswerResult(
        question_id=q.id,
        is_correct=is_correct,
        correct_answer=q.answer,
        user_answer=data.user_answer,
        analysis=q.analysis,
        recorded_as_wrong=recorded,
        score=score,
        ai_comment=ai_comment,
        sub_results=sub_results,
    )


@router.get("/admin/export/excel", summary="【管理员】导出题库Excel")
async def export_questions_excel(
    admin: User = Depends(get_current_teacher_or_admin),
    db: AsyncSession = Depends(get_db),
):
    from io import BytesIO
    from openpyxl import Workbook
    from urllib.parse import quote

    questions = (await db.execute(select(Question).order_by(Question.id.asc()))).scalars().all()

    wb = Workbook()
    ws = wb.active
    ws.title = "题库"

    # 表头
    headers = ["题目编号", "题目标题", "题目内容", "题型", "难度", "学段", "知识点", "选项", "正确答案", "解析", "创建时间"]
    ws.append(headers)

    # 列宽
    col_widths = [12, 40, 50, 10, 10, 10, 20, 40, 20, 40, 20]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w

    type_map = {"single": "单选", "multiple": "多选", "fill": "填空", "calculation": "计算"}
    diff_map = {"easy": "简单", "medium": "中等", "hard": "困难"}

    for q in questions:
        q_type = q.question_type.value if hasattr(q.question_type, 'value') else q.question_type
        q_diff = q.difficulty.value if hasattr(q.difficulty, 'value') else q.difficulty
        ws.append([
            q.question_code or f"Q{q.id:04d}",
            q.title,
            q.content or "",
            type_map.get(q_type, q_type),
            diff_map.get(q_diff, q_diff),
            q.grade or "",
            q.knowledge_point or "",
            q.options or "",
            q.answer,
            q.analysis or "",
            q.created_at.strftime("%Y-%m-%d %H:%M") if q.created_at else "",
        ])

    # 表头加粗
    from openpyxl.styles import Font
    for cell in ws[1]:
        cell.font = Font(bold=True)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = "question_bank.xlsx"
    from fastapi.responses import StreamingResponse
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
    )


@router.post("/batch-import", summary="【老师/管理员】Excel批量导入题目")
async def batch_import_questions(
    file: UploadFile = File(...),
    admin: User = Depends(get_current_teacher_or_admin),
    db: AsyncSession = Depends(get_db),
):
    """Excel批量导入题目
    Excel列：题目标题、题目内容、题型(单选/多选/填空/计算)、难度(简单/中等/困难)、学段、知识点、选项(JSON)、正确答案、解析
    """
    from io import BytesIO
    from openpyxl import load_workbook
    import json

    if not file.filename.endswith(('.xlsx', '.xls')):
        raise HTTPException(400, "请上传Excel文件(.xlsx/.xls)")

    content = await file.read()
    try:
        wb = load_workbook(BytesIO(content))
        ws = wb.active
    except Exception as e:
        raise HTTPException(400, f"Excel解析失败: {str(e)}")

    type_map = {"单选": "single", "多选": "multiple", "填空": "fill", "计算": "calculation",
                "single": "single", "multiple": "multiple", "fill": "fill", "calculation": "calculation"}
    diff_map = {"简单": "easy", "中等": "medium", "困难": "hard",
                "easy": "easy", "medium": "medium", "hard": "hard"}

    # 获取当前最大编号
    max_code = (await db.execute(
        select(Question.question_code).where(Question.question_code.isnot(None)).order_by(Question.id.desc()).limit(1)
    )).scalar()
    next_num = 1
    if max_code and max_code.startswith("Q"):
        try:
            next_num = int(max_code[1:]) + 1
        except ValueError:
            pass

    success_count = 0
    fail_count = 0
    errors = []

    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not row or not row[0]:
            continue  # 跳过空行

        try:
            title = str(row[0]).strip() if row[0] else ""
            content = str(row[1]).strip() if row[1] else ""
            q_type = type_map.get(str(row[2]).strip() if row[2] else "single", "single")
            difficulty = diff_map.get(str(row[3]).strip() if row[3] else "medium", "medium")
            grade = str(row[4]).strip() if row[4] else ""
            knowledge = str(row[5]).strip() if row[5] else ""
            options_raw = str(row[6]).strip() if row[6] else ""
            answer = str(row[7]).strip() if row[7] else ""
            analysis = str(row[8]).strip() if row[8] else ""

            if not title:
                fail_count += 1
                errors.append(f"第{row_idx}行：标题为空")
                continue
            if not answer:
                fail_count += 1
                errors.append(f"第{row_idx}行：答案为空")
                continue

            # 解析选项（存为JSON字符串）
            options = None
            if options_raw:
                try:
                    parsed = json.loads(options_raw)
                    options = json.dumps(parsed, ensure_ascii=False)  # 存为JSON字符串
                except json.JSONDecodeError:
                    options = options_raw  # 已经是字符串，直接存

            # 确保题型和难度是小写
            q_type = q_type.lower()
            difficulty = difficulty.lower()

            question_code = f"Q{next_num:04d}"
            next_num += 1

            q = Question(
                question_code=question_code,
                title=title,
                content=content,
                question_type=q_type,
                difficulty=difficulty,
                grade=grade,
                knowledge_point=knowledge,
                options=options,
                answer=answer,
                analysis=analysis,
                creator_id=admin.id,
            )
            db.add(q)
            success_count += 1
        except Exception as e:
            fail_count += 1
            errors.append(f"第{row_idx}行：{str(e)}")

    await db.commit()

    return {
        "success": success_count,
        "failed": fail_count,
        "errors": errors[:20],  # 最多返回20条错误
        "message": f"成功导入{success_count}题，失败{fail_count}题"
    }
