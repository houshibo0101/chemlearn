"""
管理员用户管理：用户列表、Excel批量导入学生、删除用户、重置密码、禁用/启用
"""
import io
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Query
from sqlalchemy import select, func
from sqlalchemy.ext.asyncio import AsyncSession
from app.core.database import get_db
from app.core.security import hash_password
from app.core.deps import get_current_admin
from app.models import User, UserRole, OfflineBind
from app.schemas import UserOut, ImportResult

router = APIRouter(prefix="/admin/users", tags=["管理员-用户管理"])


@router.get("", summary="用户列表（分页，含离线子账号）")
async def list_users(
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
    keyword: str = Query("", description="按用户名/姓名搜索"),
    role: str = Query("", description="按角色筛选：student/admin"),
    db: AsyncSession = Depends(get_db),
    admin: User = Depends(get_current_admin),
):
    stmt = select(User)
    count_stmt = select(func.count(User.id))
    if keyword:
        like = f"%{keyword}%"
        stmt = stmt.where(User.username.like(like) | User.full_name.like(like))
        count_stmt = count_stmt.where(User.username.like(like) | User.full_name.like(like))
    if role:
        stmt = stmt.where(User.role == role)
        count_stmt = count_stmt.where(User.role == role)

    total = (await db.execute(count_stmt)).scalar() or 0
    stmt = stmt.order_by(User.id.desc()).offset((page - 1) * page_size).limit(page_size)
    users = (await db.execute(stmt)).scalars().all()

    # 查询每个用户的离线子账号
    items = []
    for u in users:
        user_data = UserOut.model_validate(u).model_dump()
        # 查询绑定的离线子账号
        bind_result = await db.execute(
            select(OfflineBind).where(OfflineBind.parent_user_id == u.id)
        )
        binds = bind_result.scalars().all()
        children = []
        for bind in binds:
            child_result = await db.execute(
                select(User).where(User.id == bind.offline_user_id)
            )
            child = child_result.scalar_one_or_none()
            if child:
                child_data = UserOut.model_validate(child).model_dump()
                child_data['is_offline_child'] = True
                child_data['bound_at'] = bind.created_at.isoformat() if bind.created_at else None
                children.append(child_data)
        user_data['offline_children'] = children
        user_data['offline_child_count'] = len(children)
        items.append(user_data)

    return {
        "total": total,
        "page": page,
        "page_size": page_size,
        "items": items,
    }


@router.post("/import", response_model=ImportResult, summary="Excel批量导入学生")
async def import_students(
    file: UploadFile = File(..., description="Excel文件，需含'姓名'列，可选'学段'列"),
    db: AsyncSession = Depends(get_db),
    admin: User = Depends(get_current_admin),
):
    if not file.filename or not file.filename.lower().endswith((".xlsx", ".xls")):
        raise HTTPException(status_code=400, detail="请上传 .xlsx 或 .xls 格式的Excel文件")

    try:
        from pypinyin import lazy_pinyin
    except ImportError:
        raise HTTPException(status_code=500, detail="服务器未安装pypinyin，请执行 pip install pypinyin")

    content = await file.read()
    filename = file.filename.lower()

    # 根据文件格式选择读取方式
    rows = []
    if filename.endswith('.xlsx'):
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise HTTPException(status_code=500, detail="服务器未安装openpyxl")
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        wb.close()
    elif filename.endswith('.xls'):
        try:
            import xlrd
        except ImportError:
            raise HTTPException(status_code=500, detail="服务器未安装xlrd，请执行 pip install xlrd==1.2.0")
        wb = xlrd.open_workbook(file_contents=content)
        ws = wb.sheet_by_index(0)
        for r in range(ws.nrows):
            rows.append(tuple(ws.cell_value(r, c) for c in range(ws.ncols)))
    else:
        raise HTTPException(status_code=400, detail="仅支持 .xlsx 或 .xls 格式")

    if not rows:
        raise HTTPException(status_code=400, detail="Excel文件为空")

    # 读取表头，定位列
    headers = {}
    header_row = rows[0]
    for idx, cell in enumerate(header_row):
        if cell:
            headers[str(cell).strip()] = idx

    # 找姓名列
    name_col = None
    for key in ["姓名", "名字", "学生姓名", "name", "Name"]:
        if key in headers:
            name_col = headers[key]
            break
    if name_col is None:
        name_col = 0  # 默认第一列

    # 找学段列
    grade_col = None
    for key in ["学段", "年级", "grade", "Grade"]:
        if key in headers:
            grade_col = headers[key]
            break

    # 找班级列
    class_col = None
    for key in ["班级", "班", "class", "Class", "班级名称"]:
        if key in headers:
            class_col = headers[key]
            break

    # 预加载所有班级，避免重复查询
    from app.models import ClassRoom
    all_classes = (await db.execute(select(ClassRoom))).scalars().all()
    class_map = {c.name: c.id for c in all_classes}

    details = []
    success = 0
    failed = 0
    total = 0
    default_password = "111111"

    for row in rows[1:]:
        if not row or all(c is None for c in row):
            continue
        total += 1
        name = str(row[name_col]).strip() if row[name_col] else ""
        if not name:
            failed += 1
            details.append(f"第{total + 1}行：姓名为空，跳过")
            continue

        grade = None
        if grade_col is not None and row[grade_col]:
            grade = str(row[grade_col]).strip()

        # 处理班级
        class_id = None
        class_name = None
        if class_col is not None and row[class_col]:
            class_name = str(row[class_col]).strip()
            if class_name:
                if class_name in class_map:
                    class_id = class_map[class_name]
                else:
                    # 自动创建新班级
                    new_class = ClassRoom(name=class_name, grade=grade or "")
                    db.add(new_class)
                    await db.flush()
                    class_map[class_name] = new_class.id
                    class_id = new_class.id
                    details.append(f"📢 自动创建班级：{class_name}")

        # 生成拼音用户名
        pinyin_list = lazy_pinyin(name)
        base_username = "".join(pinyin_list).lower()
        # 过滤非字母字符
        base_username = "".join(c for c in base_username if c.isalpha())
        if not base_username:
            failed += 1
            details.append(f"第{total + 1}行：{name} 无法生成拼音用户名，跳过")
            continue

        # 处理重名：加数字后缀
        username = base_username
        suffix = 1
        while True:
            existing = (await db.execute(select(User).where(User.username == username))).scalar_one_or_none()
            if not existing:
                break
            suffix += 1
            username = f"{base_username}{suffix}"

        try:
            user = User(
                username=username,
                full_name=name,
                grade=grade,
                class_id=class_id,
                hashed_password=hash_password(default_password),
                role=UserRole.STUDENT,
                must_change_password=True,
            )
            db.add(user)
            await db.flush()
            success += 1
            note = f"（重名，用户名改为{username}）" if suffix > 1 else ""
            class_note = f"，班级：{class_name}" if class_name else ""
            details.append(f"✓ {name} → 账号：{username}，密码：{default_password}{class_note}{note}")
        except Exception as e:
            failed += 1
            details.append(f"✗ {name}：导入失败 - {str(e)}")

    await db.commit()

    return ImportResult(
        total=total,
        success=success,
        failed=failed,
        details=details,
    )


@router.delete("/{user_id}", summary="删除用户")
async def delete_user(
    user_id: int,
    db: AsyncSession = Depends(get_db),
    admin: User = Depends(get_current_admin),
):
    user = (await db.execute(select(User).where(User.id == user_id))).scalar_one_or_none()
    if not user:
        raise HTTPException(status_code=404, detail="用户不存在")
    if user.role == UserRole.ADMIN and user.username == "admin":
        raise HTTPException(status_code=400, detail="不能删除内置管理员账号")
    await db.delete(user)
    await db.commit()
    return {"message": "用户已删除"}


@router.put("/{user_id}/role", summary="修改用户角色")
async def change_user_role(
    user_id: int,
    role: str = Query(..., description="角色：student/teacher/admin"),
    db: AsyncSession = Depends(get_db),
    admin: User = Depends(get_current_admin),
):
    user = (await db.execute(select(User).where(User.id == user_id))).scalar_one_or_none()
    if not user:
        raise HTTPException(status_code=404, detail="用户不存在")
    if user.username == "admin":
        raise HTTPException(status_code=400, detail="不能修改内置管理员的角色")
    valid_roles = [r.value for r in UserRole]
    if role not in valid_roles:
        raise HTTPException(status_code=400, detail=f"无效角色，可选：{', '.join(valid_roles)}")
    old_role = user.role
    user.role = role
    await db.commit()
    role_text = {"student": "学生", "teacher": "老师", "admin": "管理员"}
    return {"message": f"已将 {user.full_name or user.username} 从{role_text.get(old_role, old_role)}改为{role_text.get(role, role)}", "role": role}


@router.post("/{user_id}/reset-password", summary="重置密码为111111")
async def reset_password(
    user_id: int,
    db: AsyncSession = Depends(get_db),
    admin: User = Depends(get_current_admin),
):
    user = (await db.execute(select(User).where(User.id == user_id))).scalar_one_or_none()
    if not user:
        raise HTTPException(status_code=404, detail="用户不存在")
    user.hashed_password = hash_password("111111")
    user.must_change_password = True
    await db.commit()
    return {"message": "密码已重置为111111，下次登录需修改密码"}


@router.post("/{user_id}/toggle", summary="禁用/启用用户")
async def toggle_user(
    user_id: int,
    db: AsyncSession = Depends(get_db),
    admin: User = Depends(get_current_admin),
):
    user = (await db.execute(select(User).where(User.id == user_id))).scalar_one_or_none()
    if not user:
        raise HTTPException(status_code=404, detail="用户不存在")
    if user.role == UserRole.ADMIN and user.username == "admin":
        raise HTTPException(status_code=400, detail="不能操作内置管理员账号")
    user.is_active = not user.is_active
    await db.commit()
    return {"message": f"用户已{'启用' if user.is_active else '禁用'}", "is_active": user.is_active}


@router.put("/{user_id}/class", summary="转移学生班级")
async def change_user_class(
    user_id: int,
    class_id: int = Query(..., description="目标班级ID，传0表示移出班级"),
    db: AsyncSession = Depends(get_db),
    admin: User = Depends(get_current_admin),
):
    from app.models import ClassRoom
    user = (await db.execute(select(User).where(User.id == user_id))).scalar_one_or_none()
    if not user:
        raise HTTPException(status_code=404, detail="用户不存在")
    if user.role != UserRole.STUDENT:
        raise HTTPException(status_code=400, detail="只能转移学生的班级")

    if class_id > 0:
        classroom = (await db.execute(select(ClassRoom).where(ClassRoom.id == class_id))).scalar_one_or_none()
        if not classroom:
            raise HTTPException(status_code=404, detail="班级不存在")
        user.class_id = class_id
        new_class_name = classroom.name
    else:
        user.class_id = None
        new_class_name = "无班级"

    await db.commit()
    return {"message": f"已将 {user.full_name or user.username} 转移到 {new_class_name}", "class_id": user.class_id}


@router.get("/stats/overview", summary="【管理员】系统概览统计（存储+Token+用户）")
async def admin_stats(
    db: AsyncSession = Depends(get_db),
    admin: User = Depends(get_current_admin),
):
    from app.services import AIService
    from app.models import Question, WrongQuestion

    # 存储统计
    storage = await AIService.get_total_storage(db)
    # Token 统计
    token_usage = await AIService.get_token_usage(db)
    # 用户统计
    total_users = (await db.execute(select(func.count(User.id)))).scalar()
    student_count = (await db.execute(select(func.count(User.id)).where(User.role == UserRole.STUDENT))).scalar()
    admin_count = (await db.execute(select(func.count(User.id)).where(User.role == UserRole.ADMIN))).scalar()
    # 题目统计
    total_questions = (await db.execute(select(func.count(Question.id)))).scalar()
    # 错题统计
    total_wrong = (await db.execute(select(func.count(WrongQuestion.id)))).scalar()

    return {
        "storage": storage,
        "token_usage": token_usage,
        "users": {
            "total": total_users,
            "students": student_count,
            "admins": admin_count,
        },
        "questions": total_questions,
        "wrong_questions": total_wrong,
    }


@router.get("/suspicious/list", summary="获取数据异常用户列表")
async def list_suspicious_users(
    db: AsyncSession = Depends(get_db),
    admin: User = Depends(get_current_admin),
):
    """获取所有被标记为数据异常的用户"""
    from app.models import APICallLog
    from app.services.rate_monitor import get_user_stats

    users = (await db.execute(
        select(User).where(User.tag == "suspicious").order_by(User.id)
    )).scalars().all()

    result = []
    for u in users:
        stats = await get_user_stats(db, u.id)
        result.append({
            "id": u.id,
            "username": u.username,
            "full_name": u.full_name,
            "email": u.email,
            "role": u.role.value if hasattr(u.role, 'value') else u.role,
            "tag": u.tag,
            "is_active": u.is_active,
            "api_stats": stats,
            "created_at": u.created_at.isoformat() if u.created_at else None,
        })
    return {"items": result, "total": len(result)}


@router.get("/{user_id}/api-stats", summary="获取用户API调用统计")
async def get_user_api_stats(
    user_id: int,
    db: AsyncSession = Depends(get_db),
    admin: User = Depends(get_current_admin),
):
    """获取指定用户的API调用统计"""
    from app.services.rate_monitor import get_user_stats
    user = (await db.execute(select(User).where(User.id == user_id))).scalar_one_or_none()
    if not user:
        raise HTTPException(status_code=404, detail="用户不存在")
    stats = await get_user_stats(db, user_id)
    return {"user_id": user_id, "username": user.username, **stats}


@router.put("/{user_id}/tag", summary="修改用户标签")
async def change_user_tag(
    user_id: int,
    tag: str = Query(..., description="标签：normal正常/suspicious数据异常/banned已封禁"),
    db: AsyncSession = Depends(get_db),
    admin: User = Depends(get_current_admin),
):
    """修改用户标签"""
    user = (await db.execute(select(User).where(User.id == user_id))).scalar_one_or_none()
    if not user:
        raise HTTPException(status_code=404, detail="用户不存在")
    if user.username == "admin":
        raise HTTPException(status_code=400, detail="不能修改内置管理员的标签")
    valid_tags = ["normal", "suspicious", "banned"]
    if tag not in valid_tags:
        raise HTTPException(status_code=400, detail=f"无效标签，可选：{', '.join(valid_tags)}")
    old_tag = user.tag
    user.tag = tag
    # 如果标记为banned，同时禁用账号
    if tag == "banned":
        user.is_active = False
    elif tag == "normal" and not user.is_active:
        user.is_active = True
    await db.commit()
    tag_text = {"normal": "正常", "suspicious": "数据异常", "banned": "已封禁"}
    return {"message": f"已将 {user.full_name or user.username} 从{tag_text.get(old_tag, old_tag)}标记为{tag_text.get(tag, tag)}", "tag": tag}


@router.post("/{user_id}/ban", summary="封禁用户")
async def ban_user(
    user_id: int,
    reason: str = Query("", description="封禁原因"),
    db: AsyncSession = Depends(get_db),
    admin: User = Depends(get_current_admin),
):
    """封禁用户"""
    user = (await db.execute(select(User).where(User.id == user_id))).scalar_one_or_none()
    if not user:
        raise HTTPException(status_code=404, detail="用户不存在")
    if user.username == "admin":
        raise HTTPException(status_code=400, detail="不能封禁内置管理员账号")
    user.is_active = False
    user.tag = "banned"
    await db.commit()
    return {"message": f"已封禁用户 {user.full_name or user.username}", "is_active": False}


@router.post("/{user_id}/unban", summary="解封用户")
async def unban_user(
    user_id: int,
    db: AsyncSession = Depends(get_db),
    admin: User = Depends(get_current_admin),
):
    """解封用户"""
    user = (await db.execute(select(User).where(User.id == user_id))).scalar_one_or_none()
    if not user:
        raise HTTPException(status_code=404, detail="用户不存在")
    user.is_active = True
    user.tag = "normal"
    await db.commit()
    return {"message": f"已解封用户 {user.full_name or user.username}", "is_active": True}


# ========== 管理员管理API Key ==========

@router.get("/{user_id}/api-keys", summary="【管理员】查看用户的API Key列表（不显示完整key）")
async def admin_list_user_api_keys(
    user_id: int,
    db: AsyncSession = Depends(get_db),
    admin: User = Depends(get_current_admin),
):
    """查看指定用户的API Key列表，只显示前缀，不显示完整key"""
    from app.models import APIKey
    from app.schemas.api_key import api_key_to_out

    user = (await db.execute(select(User).where(User.id == user_id))).scalar_one_or_none()
    if not user:
        raise HTTPException(status_code=404, detail="用户不存在")

    result = await db.execute(
        select(APIKey).where(APIKey.user_id == user_id).order_by(APIKey.created_at.desc())
    )
    items = result.scalars().all()
    return {
        "user_id": user_id,
        "username": user.username,
        "items": [api_key_to_out(i) for i in items],
        "total": len(items),
    }


@router.put("/api-keys/{key_id}/limit", summary="【管理员】修改API Key限额")
async def admin_update_api_key_limit(
    key_id: int,
    daily_limit: int = Query(None, ge=0, description="每日调用限额，0表示不限"),
    rate_limit_per_minute: int = Query(None, ge=0, description="每分钟调用限额，0表示不限"),
    is_active: bool = Query(None, description="是否启用"),
    db: AsyncSession = Depends(get_db),
    admin: User = Depends(get_current_admin),
):
    """管理员修改API Key的调用限额和启用状态，看不到完整key"""
    from app.models import APIKey

    key_obj = (await db.execute(select(APIKey).where(APIKey.id == key_id))).scalar_one_or_none()
    if not key_obj:
        raise HTTPException(status_code=404, detail="API Key不存在")

    changes = []
    if daily_limit is not None:
        key_obj.daily_limit = daily_limit
        changes.append(f"每日限额={daily_limit}")
    if rate_limit_per_minute is not None:
        key_obj.rate_limit_per_minute = rate_limit_per_minute
        changes.append(f"每分钟限额={rate_limit_per_minute}")
    if is_active is not None:
        key_obj.is_active = is_active
        changes.append(f"{'启用' if is_active else '禁用'}")

    await db.commit()

    return {
        "message": f"API Key已更新：{', '.join(changes)}",
        "key_id": key_id,
        "key_prefix": key_obj.key[:8] + "..." if len(key_obj.key) > 8 else key_obj.key,
        "daily_limit": key_obj.daily_limit,
        "rate_limit_per_minute": key_obj.rate_limit_per_minute,
        "is_active": key_obj.is_active,
    }


@router.delete("/api-keys/{key_id}", summary="【管理员】删除用户的API Key")
async def admin_delete_api_key(
    key_id: int,
    db: AsyncSession = Depends(get_db),
    admin: User = Depends(get_current_admin),
):
    """管理员删除指定API Key"""
    from app.models import APIKey

    key_obj = (await db.execute(select(APIKey).where(APIKey.id == key_id))).scalar_one_or_none()
    if not key_obj:
        raise HTTPException(status_code=404, detail="API Key不存在")

    key_prefix = key_obj.key[:8] + "..." if len(key_obj.key) > 8 else key_obj.key
    user_id = key_obj.user_id

    await db.delete(key_obj)
    await db.commit()

    return {
        "message": f"API Key {key_prefix} 已删除",
        "key_id": key_id,
        "user_id": user_id,
    }
